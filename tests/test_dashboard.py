"""Browser tests for the static dashboard (Playwright + Chrome)."""
import re

CHARTS = ("share", "change", "shareQ", "channelQ", "digital", "digitalQ", "regions", "regionsEbit", "bridge", "marginsQ",
          "inventoryQ", "daysQ")

ECHARTS_SERIES = """id => {
  const c = echarts.getInstanceByDom(document.querySelector('.chart[data-chart="' + id + '"]'));
  const o = c.getOption();
  const value = d => d && typeof d === 'object' && !Array.isArray(d) ? d.value : d;
  return {x: o.xAxis[0].data || null, y: o.yAxis[0].data || null,
          series: o.series.map(s => ({name: s.name || null, type: s.type, data: (s.data || []).map(value)}))};
}"""

# Cell texts; a row label without the unit it repeats for phones (span.unit-inline).
TABLE = """() => Array.from(document.querySelectorAll('#data-table tbody tr')).filter(r => !r.classList.contains('group'))
    .map(r => Array.from(r.children).map(c => Array.from(c.childNodes)
        .filter(n => !(n.classList && n.classList.contains('unit-inline'))).map(n => n.textContent).join('').trim()))"""


def vi_number(text):
    """Parse a vi-VN formatted figure from a table cell: '12.382', '(50)', '−4,4%', '+8%', '≈1,7' (page reference after)."""
    t = text.split("tr.")[0].replace("†", "").replace("%", "").replace("≈", "").strip()
    neg = t.startswith("(") or t.startswith("−") or t.startswith("-")
    t = t.strip("()+−-").replace(".", "").replace(",", ".")
    v = float(t)
    return -v if neg else v


def show_table(page, kind):
    page.check(f'[data-toggle="table"] input[value="{kind}"]', force=True)


def test_initial_render_has_no_errors(page, site_data):
    page.open()
    assert page.console_errors == []
    for chart in CHARTS:
        assert page.locator(f'.chart[data-chart="{chart}"] canvas').count() == 1, chart
        assert "Số liệu chi tiết" in page.get_attribute(f'.chart[data-chart="{chart}"]', "aria-label"), chart
    assert page.locator("#findings .finding").count() == len(site_data["findings"]) == 6
    assert page.locator("#quotes .quote").count() == len(site_data["quotes"])
    assert page.locator("#data-table thead th").count() == 2 + len(site_data["periods"]["A"])


def test_findings_show_the_pipeline_text(page, site_data):
    page.open()
    cards = page.eval_on_selector_all("#findings .finding", """fs => fs.map(f => ({
        label: f.querySelector('h3').textContent, value: f.querySelector('.finding-value').textContent,
        text: f.querySelector('.finding-text').textContent, href: f.querySelector('a').getAttribute('href')}))""")
    assert cards == [{"label": f["label"], "value": f["value"], "text": f["text"], "href": "#" + f["target"]}
                     for f in site_data["findings"]]


def test_annual_table_matches_data(page, site_data):
    page.open()
    periods = site_data["periods"]["A"]
    rows = {r[0]: r for r in page.evaluate(TABLE)}
    for label, metric in (("NIKE Direct", "nike_direct_revenue"), ("Doanh thu", "revenues"), ("Hàng tồn kho cuối kỳ", "inventories"),
                          ("EBIT do NIKE công bố", "ebit_reported"), ("Chi phí (thu nhập) lãi thuần", "interest_expense_net")):
        for i, p in enumerate(periods):
            assert vi_number(rows[label][2 + i]) == site_data["observations"][metric][p]["value"], (label, p)
    for i, p in enumerate(periods):
        share = round(site_data["kpis"]["direct_share"]["values"][p] * 100, 1)
        assert round(vi_number(rows["Tỷ trọng NIKE Direct"][2 + i]), 1) == share


def test_quarterly_table(page, site_data):
    page.open()
    show_table(page, "Q")
    periods = site_data["periods"]["Q"]
    assert page.locator("#data-table thead tr.years th.year").all_inner_texts() == site_data["periods"]["A"]
    assert page.locator("#data-table thead tr:last-child th.period").all_inner_texts() == ["Q1", "Q2", "Q3", "Q4"] * 3
    rows = {r[0]: r for r in page.evaluate(TABLE)}
    for i, p in enumerate(periods):
        assert vi_number(rows["NIKE Direct"][2 + i]) == site_data["observations"]["nike_direct_revenue"][p]["value"], p
    # A fourth-quarter flow is calculated, so it has no page reference; only the 31 May balance (10-K page) and the
    # percentage stated in the fourth-quarter earnings release ("8-K") link to a document. An estimate from rounded
    # figures keeps one decimal.
    q4 = periods.index("FY2020Q4")
    assert rows["Doanh số Digital"][2 + q4] == "≈1,7"
    assert rows["Tỷ trọng NIKE Direct do NIKE nêu (xấp xỉ)"][2 + q4] == "Không công bố"
    assert rows["Tăng trưởng doanh thu, loại trừ tỷ giá (non-GAAP)"][2 + q4] == "−36%8-K"
    for i in (j for j, p in enumerate(periods) if p.endswith("Q4")):
        linked = page.eval_on_selector_all("#data-table tbody tr:not(.group)",
                                           f"rs => rs.filter(r => r.children[{2 + i}].querySelector('a')).map(r => "
                                           "r.children[0].textContent.replace(r.children[0].querySelector('.unit-inline').textContent, ''))")
        assert linked == ["Tăng trưởng doanh thu, loại trừ tỷ giá (non-GAAP)", "Hàng tồn kho cuối kỳ"], periods[i]


def test_region_table(page, site_data):
    page.open()
    show_table(page, "R")
    regions = site_data["regions"]
    assert page.locator("#data-table tbody tr.group").all_inner_texts() == [
        r["vi"] if r["vi"] == r["short"] else f"{r['vi']} ({r['short']})" for r in regions]
    body, periods = page.evaluate(TABLE), site_data["periods"]["A"]
    size = len(body) // len(regions)
    for n, r in enumerate(regions):
        block = {row[0]: row for row in body[n * size:(n + 1) * size]}
        for label, metric in (("Doanh thu", "revenue"), ("NIKE Direct", "direct_revenue"), ("EBIT", "ebit")):
            for i, p in enumerate(periods):
                assert vi_number(block[label][2 + i]) == site_data["observations"][f"{r['id']}_{metric}"][p]["value"], (r["id"], label, p)


def test_charts_plot_the_data(page, site_data):
    page.open()
    A, Q = site_data["periods"]["A"], site_data["periods"]["Q"]
    k = lambda kid, ps: [site_data["kpis"][kid]["values"][p] for p in ps]
    bn = lambda mid, ps: [site_data["observations"][mid][p]["value"] / 1000 for p in ps]
    share = page.evaluate(ECHARTS_SERIES, "share")
    assert [s["name"] for s in share["series"]] == ["NIKE Direct", "Bán buôn", "Global Brand Divisions"]
    assert share["series"][0]["data"] == k("direct_share", A)
    share_q = page.evaluate(ECHARTS_SERIES, "shareQ")
    assert share_q["x"] == Q and share_q["series"][0]["data"] == k("direct_share", Q)
    channel = page.evaluate(ECHARTS_SERIES, "channelQ")
    assert [(s["name"], s["data"]) for s in channel["series"] if s["type"] == "line"] == [
        ("NIKE Direct", bn("nike_direct_revenue", Q)), ("Bán buôn", bn("nike_brand_wholesale_revenue", Q))]
    digital = page.evaluate(ECHARTS_SERIES, "digital")
    assert digital["series"][0]["data"] == [site_data["observations"]["nike_brand_digital_sales"][p]["value"] for p in A]
    assert digital["series"][1]["data"] == [v / 1000 for v in k("direct_ex_digital", A)]
    assert page.evaluate(ECHARTS_SERIES, "digitalQ")["series"][0]["data"] == k("digital_share_of_direct", Q)
    margins = page.evaluate(ECHARTS_SERIES, "marginsQ")["series"]
    assert [s["data"] for s in margins if s["type"] == "line"] == [k("gross_margin", Q), k("ebit_margin", Q)]
    assert page.evaluate(ECHARTS_SERIES, "inventoryQ")["series"][0]["data"] == bn("inventories", Q)
    assert page.evaluate(ECHARTS_SERIES, "daysQ")["series"][0]["data"] == k("days_inventory", Q)
    # Regions: a FY2019 dot and a FY2021 dot per row, NIKE Brand as a whole first.
    regions = page.evaluate(ECHARTS_SERIES, "regions")
    assert regions["y"] == ["Toàn NIKE Brand"] + [r["short"] for r in site_data["regions"]]
    ids = ["direct_share"] + [r["id"] + "_direct_share" for r in site_data["regions"]]
    dots = {s["name"]: [d[0] for d in s["data"]] for s in regions["series"] if s["type"] == "scatter"}
    assert dots == {y: [site_data["kpis"][i]["values"][y] for i in ids] for y in ("FY2019", "FY2021")}


COVID_BAND = """id => {
  const el = document.querySelector('.chart[data-chart="' + id + '"]'), chart = echarts.getInstanceByDom(el);
  const rect = chart.getModel().getComponent('grid').coordinateSystem.getRect();
  const x = chart.convertToPixel({xAxisIndex: 0}, 'FY2020Q4'), step = chart.convertToPixel({xAxisIndex: 0}, 'FY2021Q1') - x;
  const canvas = el.querySelector('canvas'), ctx = canvas.getContext('2d');
  const dpr = canvas.width / canvas.getBoundingClientRect().width;   // the chart draws at 2x or more
  // Opacity of the canvas just below the top of the plot: the band is the only thing drawn there.
  const alpha = dx => ctx.getImageData(Math.round((x + dx) * dpr), Math.round((rect.y + 2) * dpr), 1, 1).data[3];
  return {inside: [alpha(-0.4 * step), alpha(0), alpha(0.4 * step)], outside: [alpha(-0.6 * step), alpha(0.6 * step)]};
}"""


def test_fourth_quarter_of_fy2020_is_shaded(page):
    """The quarterly charts shade the whole width of Q4 FY2020 (March–May 2020), and nothing beside it."""
    page.open()
    for chart in ("shareQ", "channelQ", "digitalQ", "marginsQ", "inventoryQ", "daysQ"):
        shade = page.evaluate(COVID_BAND, chart)
        assert all(a > 0 for a in shade["inside"]) and shade["outside"] == [0, 0], (chart, shade)


def test_period_toggles_redraw_the_comparisons(page, site_data):
    page.open()
    k = lambda kid, p: site_data["kpis"][kid]["values"][p]
    v = lambda mid, p: site_data["observations"][mid][p]["value"]
    channels = ["nike_direct_revenue", "nike_brand_wholesale_revenue", "global_brand_divisions_revenue", "nike_brand_revenue"]
    for a, b in (("FY2019", "FY2021"), ("FY2019", "FY2020"), ("FY2020", "FY2021")):
        page.check(f'[data-toggle="change"] input[value="{a}-{b}"]', force=True)
        page.check(f'[data-toggle="bridge"] input[value="{a}-{b}"]', force=True)
        assert page.inner_text("#change-sub") == f"Tỷ USD, thay đổi từ {a} đến {b}. Số âm là giảm."
        assert page.evaluate(ECHARTS_SERIES, "change")["series"][0]["data"] == [(v(m, b) - v(m, a)) / 1000 for m in channels]
        bars = page.evaluate(ECHARTS_SERIES, "bridge")["series"][1]["data"]
        assert bars[0] == k("ebit_margin", a) and bars[-1] == k("ebit_margin", b), (a, b)
        assert f"biên EBIT năm {a} và {b}" in page.inner_text("#bridge-sub")
    assert page.console_errors == []


def test_tooltip_shows_units_and_source(page):
    page.open()
    chart = page.locator('.chart[data-chart="channelQ"]')
    chart.scroll_into_view_if_needed()
    box = chart.bounding_box()
    page.mouse.move(box["x"] + box["width"] * 0.6, box["y"] + box["height"] * 0.5)
    page.wait_for_timeout(400)
    tip = page.evaluate("""() => Array.from(document.querySelectorAll('.chart[data-chart="channelQ"] div'))
        .map(d => d.innerText).filter(t => t.includes('tỷ USD')).join(' ')""")
    assert "tỷ USD" in tip and "so với cùng kỳ" in tip
    assert re.search(r"Nguồn: 10-Q Q[1-3] FY20\d\d, tr\. \d+|Quý 4 tự tính: cả năm ở 10-K FY20\d\d", tip), tip


def test_source_links_and_footer(page, site_data, server):
    page.open()
    allowed = {f["url"] for f in site_data["filings"].values()}
    for kind in ("A", "Q", "R"):
        show_table(page, kind)
        hrefs = page.eval_on_selector_all("#data-table a.page-ref", "as => as.map(a => a.href)")
        assert hrefs and set(hrefs) <= allowed, kind
    assert all(h.startswith("https://www.sec.gov/Archives/edgar/data/320187/") for h in allowed)
    # Annual reports, quarterly reports, fourth-quarter earnings releases, and the later reports used only to look for
    # restatements.
    groups = page.eval_on_selector_all("#filing-list tbody", "bs => bs.map(b => b.querySelectorAll('a').length)")
    assert groups == [3, 9, 3, 5] and sum(groups) == len(site_data["filings"])
    href = page.get_attribute("#evidence-link", "href")
    assert href == site_data["meta"]["evidence"] and page.request.get(server + href).ok
    total = sum(site_data["meta"]["status_counts"].values())
    assert page.inner_text("#status-summary").startswith(f"{total} số liệu: ")


def test_xlsx_and_png_exports(page, site_data, tmp_path):
    from openpyxl import load_workbook
    page.open()
    for kind, name in (("A", "theo-nam"), ("Q", "theo-quy")):
        periods = site_data["periods"][kind]
        show_table(page, kind)
        with page.expect_download() as d:
            page.click("#xlsx-table")
        assert d.value.suggested_filename == f"NIKE_bang-{name}_FY2019-FY2021.xlsx"
        saved = tmp_path / d.value.suggested_filename
        d.value.save_as(saved)
        rows = [[c.value for c in r] for r in load_workbook(saved).active.iter_rows()]
        assert rows[0][:3 + len(periods)] == ["Chỉ tiêu", "Đơn vị", "Loại"] + [site_data["labels"][p] for p in periods]
        direct = next(r for r in rows if r[0] == "NIKE Direct")
        assert direct[3:3 + len(periods)] == [site_data["observations"]["nike_direct_revenue"][p]["value"] for p in periods]
    stated = next(r for r in rows if r[0] == "Tỷ trọng NIKE Direct do NIKE nêu (xấp xỉ)")
    assert stated[3 + site_data["periods"]["Q"].index("FY2020Q4")] == "Không công bố"
    with page.expect_download() as d:
        page.click('.export-png[data-chart="shareQ"]')
    data = open(d.value.path(), "rb").read()
    assert data[:8] == b"\x89PNG\r\n\x1a\n" and len(data) > 20000
    assert d.value.suggested_filename == "NIKE_shareQ.png"


def test_workbook_download_available(page, server):
    page.open()
    href = page.get_attribute("#download-workbook", "href")
    resp = page.request.get(server + href)
    assert resp.ok and resp.body()[:2] == b"PK"


def test_accessible_names_and_headings(page):
    page.open()
    unnamed = page.evaluate("""() => Array.from(document.querySelectorAll('button, select, a, input'))
        .filter(e => e.offsetParent !== null || e.type === 'radio')
        .filter(e => !(e.getAttribute('aria-label') || e.textContent.trim() || (e.labels && e.labels.length)))
        .map(e => e.outerHTML.slice(0, 80))""")
    assert unnamed == []
    levels = page.eval_on_selector_all("h1, h2, h3", "hs => hs.filter(h => h.offsetParent !== null).map(h => +h.tagName[1])")
    assert levels[0] == 1 and all(b - a <= 1 for a, b in zip(levels, levels[1:]))
    assert page.locator("#data-table thead th[scope=col]").count() == 2 + 3
    targets = page.eval_on_selector_all(".section-nav a", "as => as.map(a => a.getAttribute('href'))")
    assert len(targets) == 8 and all(page.locator(t).count() == 1 for t in targets)


HEADER_VS_FIGURES = """() => {
  const right = n => { const r = document.createRange(); r.selectNodeContents(n); return r.getBoundingClientRect().right; };
  const heads = Array.from(document.querySelectorAll('#data-table thead tr:last-child th')).slice(2);
  return Array.from(document.querySelectorAll('#data-table tbody tr')).filter(r => !r.classList.contains('group'))
    .flatMap(r => Array.from(r.querySelectorAll('td.num')).map((td, i) =>
      [r.querySelector('th').textContent + ' / ' + heads[i].textContent, Math.abs(right(heads[i]) - right(td.firstChild))]));
}"""

# Kept on one line wherever they appear (app.js, keepTogether): a report with its period, "non-GAAP", a period range, a
# figure with its unit, "(xấp xỉ)" and "(ước tính)".
UNBROKEN = r"""() => {
  const bad = [];
  document.querySelectorAll('.nowrap').forEach(s => { if (s.getClientRects().length > 1) bad.push(s.textContent); });
  const w = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
  for (let n; (n = w.nextNode());) {
    if (/10-[KQ]|non-GAAP|Q[1-4] FY\d{4}|FY\d{4}(–| → )FY\d{4}|\d (tỷ|triệu) USD|\d điểm %|\d ngày|\((xấp xỉ|ước tính)\)/.test(n.nodeValue) &&
        !n.parentNode.closest('.nowrap, script, style, option, .chart, .visually-hidden, .finding-value'))
      bad.push(n.nodeValue.trim().slice(0, 50));
  }
  return bad;
}"""

PINNED = """() => {
  const box = document.querySelector('.table-scroll'), b = box.getBoundingClientRect();
  const heads = Array.from(document.querySelectorAll('#data-table thead tr:last-child th')), edge = heads[0].getBoundingClientRect().right;
  const label = document.querySelector('#data-table tbody tr:not(.group) th').getBoundingClientRect();
  const clipped = Array.from(box.querySelectorAll('.group-label')).map(g => [g.textContent, g.getBoundingClientRect()])
    .filter(([, r]) => r.left < b.left - 0.5 || r.right > b.right + 0.5).map(([t]) => t);
  // A column cut by the pinned column's edge would show the tail of an earlier period beside the labels.
  const partial = heads.slice(1).map(h => [h.textContent, h.getBoundingClientRect()])
    .filter(([, r]) => r.left < edge - 0.5 && r.right > edge + 0.5).map(([t]) => t);
  return {overflow: box.scrollWidth > box.clientWidth, scrolled: box.scrollLeft > 0 && box.classList.contains('scrolled'),
          header: Math.round(heads[0].getBoundingClientRect().left - label.left), clipped, partial};
}"""


def assert_headers_align(pg, kind):
    """Every cell (reported, calculated or 'Không công bố') ends exactly under its period header."""
    offsets = pg.evaluate(HEADER_VS_FIGURES)
    off = [(cell, round(d, 1)) for cell, d in offsets if d > 1]
    assert offsets and not off, (kind, off)


def test_tables_fit_and_headers_align_on_desktop(page):
    page.open()
    for kind in ("A", "Q", "R"):
        show_table(page, kind)
        if kind != "Q":
            width, client = page.evaluate("() => { const s = document.querySelector('.table-scroll'); return [s.scrollWidth, s.clientWidth]; }")
            assert width <= client, (kind, width, client)
        assert_headers_align(page, kind)
    assert page.evaluate(UNBROKEN) == []
    # Charts redraw on resize, and the period toggles rewrite their subtitles: still nothing split.
    page.set_viewport_size({"width": 1180, "height": 1000})
    page.wait_for_timeout(500)
    page.check('[data-toggle="bridge"] input[value="FY2020-FY2021"]', force=True)
    assert page.evaluate(UNBROKEN) == []


def test_narrow_layouts(browser, server):
    for width in (390, 1024):
        ctx = browser.new_context(viewport={"width": width, "height": 844})
        pg = ctx.new_page()
        pg.goto(server)
        pg.wait_for_selector('html[data-ready="true"]')
        for kind in ("A", "Q", "R"):
            show_table(pg, kind)
            assert pg.evaluate("document.documentElement.scrollWidth <= document.documentElement.clientWidth"), (width, kind)
            assert_headers_align(pg, kind)
            # A table wider than the screen opens on the latest periods; the label column, its header and the group
            # names stay in view, and the first period in view starts at the label column's edge.
            pinned = pg.evaluate(PINNED)
            if pinned.pop("overflow"):
                assert pinned == {"scrolled": True, "header": 0, "clipped": [], "partial": []}, (width, kind, pinned)
        assert pg.evaluate(UNBROKEN) == [], width
        ctx.close()


def test_print_draws_every_chart(page):
    page.open()
    page.evaluate("() => window.dispatchEvent(new Event('beforeprint'))")
    page.emulate_media(media="print")
    for chart in CHARTS:
        size = page.evaluate("id => { const c = document.querySelector('.chart[data-chart=\"' + id + '\"] canvas'); "
                             "return c ? c.width * c.height : 0; }", chart)
        assert size > 0, chart
    # Paper gets all three tables, not the one the page happens to show.
    printed = page.evaluate("""() => Array.from(document.querySelectorAll('#print-tables table')).map(t =>
        [t.getBoundingClientRect().height > 0, t.querySelectorAll('tbody tr').length])""")
    assert len(printed) == 3 and all(shown and rows > 20 for shown, rows in printed), printed
    assert page.evaluate("document.querySelector('#bang-so-lieu .table-scroll').getBoundingClientRect().height") == 0
    page.emulate_media(media="screen")
    assert page.evaluate("document.getElementById('print-tables').getBoundingClientRect().height") == 0
    page.evaluate("() => window.dispatchEvent(new Event('afterprint'))")


def test_reduced_motion_disables_chart_animation(browser, server):
    ctx = browser.new_context(reduced_motion="reduce")
    pg = ctx.new_page()
    pg.goto(server)
    pg.wait_for_selector('html[data-ready="true"]')
    animated = pg.evaluate("() => Array.from(document.querySelectorAll('.chart')).map(el => echarts.getInstanceByDom(el).getOption().animation)")
    assert len(animated) == len(CHARTS) and not any(animated)
    ctx.close()
