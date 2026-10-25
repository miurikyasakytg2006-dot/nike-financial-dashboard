"""Data-integrity tests: extraction, observations, workbook, validation, KPI formulas and the three dashboard versions."""
import json
import posixpath
import re
import shutil
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from pipeline import dataset
from pipeline.build_site_data import SITE_DOWNLOAD, compare_with_workbook_formulas, read_observations, validate
from pipeline.build_workbook import FIELD_COL, HEADER_ROW, SHEET_A, SHEET_Q, SHEET_R, WORKBOOK, git_user_name
from pipeline.insights import findings, section_notes
from pipeline.metrics import (A, CHECKS, COMPARISONS, Q, REPORTED, STATUS_CONFLICT, STATUS_HUMAN, STATUS_MATCHED,
                              STATUS_MISSING, applies)

# Representative values re-read from the filings: years, quarters (a fourth quarter is the year minus nine months),
# a balance, a region and MD&A figures.
GOLDEN = {
    ("FY2019", "revenues"): 39117, ("FY2020", "revenues"): 37403, ("FY2021", "revenues"): 44538,
    ("FY2019", "nike_direct_revenue"): 11753, ("FY2020", "nike_direct_revenue"): 12382,
    ("FY2021", "nike_direct_revenue"): 16370, ("FY2021", "nike_brand_wholesale_revenue"): 25898,
    ("FY2019", "nike_brand_revenue"): 37218, ("FY2020", "inventories"): 7367, ("FY2021", "ebit_reported"): 6923,
    ("FY2021", "interest_expense_net"): 262, ("FY2021", "corporate_revenue"): 40, ("FY2021", "gc_ebit"): 3243,
    ("FY2021", "nike_brand_digital_sales"): Decimal("9.1"), ("FY2019Q1", "revenues"): 9948,
    ("FY2019Q1", "nike_brand_digital_sales"): Decimal("0.77"), ("FY2020Q4", "revenues"): 6313,
    ("FY2020Q4", "net_income"): -790, ("FY2020Q4", "inventories"): 7367, ("FY2021Q4", "revenues"): 12344,
}
PAGES = ["tong_quan", "theo_quy", "khu_vuc_loi_nhuan", "bang_so_lieu", "nguon"]


@pytest.fixture(scope="module")
def facts():
    return dataset.load_facts()


@pytest.fixture(scope="module")
def assembled(facts):
    return dataset.assemble(facts)


def test_every_metric_and_period_present_once(assembled):
    obs, _ = assembled
    keys = [o["key"] for o in obs]
    assert len(keys) == len(set(keys)) == sum(len(dataset.PERIODS[f]) for m in REPORTED for f in (A, Q) if applies(m, f))
    # Only the percentages NIKE states are missing, and only for the fourth quarter, which has no 10-Q.
    missing = [o for o in obs if o["value"] is None]
    percentages = {m["id"] for m in REPORTED if m["unit"] == "%"}
    assert {o["metric"]["id"] for o in missing} == percentages and len(missing) == 3 * len(percentages)
    assert all(o["quarter"] == 4 and o["status"] == STATUS_MISSING for o in missing)


@pytest.mark.parametrize("key,value", sorted(GOLDEN.items()))
def test_golden_values(assembled, key, value):
    obs, _ = assembled
    o = next(o for o in obs if (o["period"], o["metric"]["id"]) == key)
    assert o["value"] == value


def test_fourth_quarter_is_the_year_minus_nine_months(facts, assembled):
    obs, _ = assembled
    by = {(o["period"], o["metric"]["id"]): o for o in obs}
    checked = 0
    for o in obs:
        if o["quarter"] != 4 or o["value"] is None:
            continue
        m, fy = o["metric"], o["fy"]
        year = by[(f"FY{fy}", m["id"])]["value"]
        if dataset.is_instant(m):      # a balance: 31 May, the same date as the year
            assert o["value"] == year, o["key"]
            continue
        nine = facts[(f"10-Q Q3 FY{fy}", m["id"], f"FY{fy}M9", dataset.method_of(m))]["value"]
        assert o["method"] == "q4_derived" and o["value"] == year - nine, o["key"]
        # The four quarters add up to the year, give or take the rounding of the nine-month figure.
        quarters = sum(by[(f"FY{fy}Q{q}", m["id"])]["value"] for q in (1, 2, 3, 4))
        assert abs(quarters - year) <= dataset.rounding_tolerance(m["unit"], 3), o["key"]
        checked += 1
    assert checked > 50


def test_quarters_add_up_to_the_year_to_date_figures(facts):
    checks = dataset.ytd_checks(facts)
    assert len(checks) > 150 and [c for c in checks if not c["ok"]] == []


def test_every_figure_has_a_second_source(assembled):
    """Each figure is confirmed by an independent presentation (MD&A table, notes, year-to-date sums, the quarterly
    table of the 10-K, a later report, a reading of the page); nothing is pending or conflicting."""
    obs, _ = assembled
    for o in obs:
        if o["value"] is not None:
            assert o["status"] == STATUS_MATCHED and o["confirms"] and not o["conflicts"], (o["key"], o["cross_check"])


def test_later_reports_do_not_change_the_figures(assembled):
    obs, _ = assembled
    assert [o["key"] for o in obs if o["restatement"].startswith("CÓ THAY ĐỔI")] == []


def test_reconciliations_and_comparisons_within_tolerance(assembled):
    _, results = assembled
    for period, env in results.items():
        for c in CHECKS + COMPARISONS:
            v = env[c["id"]]
            if c in CHECKS and applies(c, dataset.freq_of(period)):
                assert v is not None, (c["id"], period)
            if v is not None:
                assert abs(v) <= Decimal(str(c["tol"])), (c["id"], period, v)


def test_kpi_formulas_against_hand_calculation(assembled):
    _, r = assembled
    assert r["FY2020"]["gross_margin"] == Decimal(16241) / Decimal(37403)
    assert r["FY2021"]["direct_share"] == Decimal(16370) / Decimal(42293)
    assert r["FY2021"]["direct_growth"] == Decimal(16370) / Decimal(12382) - 1
    assert r["FY2021"]["ebit"] == 5727 + 934 + 262
    assert r["FY2020"]["average_inventory"] == (Decimal(7367) + Decimal(5622)) / 2
    # A quarter: a quarter of a year, the quarter's opening and closing balances, its own cost of sales.
    assert r["FY2020Q4"]["days_inventory"] == Decimal("91.25") * ((Decimal(7367) + Decimal(5807)) / 2) / Decimal(3960)
    assert r["FY2020Q4"]["direct_growth"] == Decimal(2581) / Decimal(3105) - 1


def test_missing_values_never_become_zero(assembled):
    _, r = assembled
    assert r["FY2020Q4"]["cmp_direct_share"] is None
    assert r["FY2019Q1"]["inventory_growth"] is None      # the same quarter of FY2018 is outside the scope
    assert dataset.evaluate("a / b", {"a": Decimal(1), "b": None}) is None


def test_quarterly_mdna_figures_were_read_again_on_the_page(assembled):
    """Currency-neutral growth of a quarter is printed only in that 10-Q's MD&A and never again, so each figure was
    read a second time on the rendered page, with a screenshot."""
    from pipeline.recheck_pages import EVIDENCE_PAGE
    obs, _ = assembled
    checks = dataset.load_page_check()
    targets = [o for o in obs if o["metric"]["id"].endswith("_pct_change_cn") and o["freq"] == Q and o["quarter"] < 4]
    assert len(targets) == len(checks) == 36
    for o in targets:
        c = checks[(o["source"], o["metric"]["id"], o["period"])]
        assert c["result"] == "match" and Decimal(c["extracted_value"]) == o["value"], o["key"]
        assert "đã đọc lại ngay trên trang" in o["cross_check"].lower(), o["key"]
        assert (dataset.ROOT / "handover" / "bang-chung-doi-chieu" / c["screenshot"]).exists(), c["screenshot"]
    for folder in ("handover", "site"):
        assert (dataset.ROOT / folder / "bang-chung-doi-chieu" / EVIDENCE_PAGE).exists(), folder


# ---------------------------------------------------------------- workbook validation
def test_workbook_passes_validation(assembled):
    obs, _ = assembled
    assert validate(read_observations(WORKBOOK), obs) == []


def _mutated_copy(tmp_path, sheet, key, field, new_value):
    path = tmp_path / "wb.xlsx"
    shutil.copyfile(WORKBOOK, path)
    wb = load_workbook(path)
    for row in wb[sheet].iter_rows(min_row=HEADER_ROW + 2):
        if row[FIELD_COL["key"]].value == key:
            row[FIELD_COL[field]].value = new_value
            break
    else:
        raise AssertionError(key)
    wb.save(path)
    return path


@pytest.mark.parametrize("sheet,key,field,value,expected", [
    (SHEET_A, "FY2021|revenues", "value", 45000, "differs from the extraction"),
    (SHEET_Q, "Q4 FY2020|revenues", "value", 6000, "differs from the extraction"),
    (SHEET_Q, "Q2 FY2020|nike_direct_revenue", "value", None, "must be marked missing"),
    (SHEET_Q, "Q3 FY2021|nike_direct_pct_change_cn", "value", 16, "whole percent"),
    (SHEET_A, "FY2020|cost_of_sales", "unit", "tỷ USD", "unit"),
    (SHEET_R, "FY2021|na_revenue", "scope", "NIKE, Inc.", "reporting scope"),
    (SHEET_A, "FY2019|gross_profit", "status", STATUS_CONFLICT, "conflicting evidence"),
    (SHEET_Q, "Q1 FY2021|nike_direct_revenue", "status", STATUS_HUMAN, "verifier's name"),
    (SHEET_A, "FY2020|net_income", "page", None, "incomplete provenance"),
    (SHEET_A, "FY2020|net_income", "page", 99, "page 99 differs"),
    (SHEET_A, "FY2021|inventories", "prior", 6000, "comparative"),
    (SHEET_R, "FY2021|gc_revenue", "key", "FY2021|na_revenue", "Duplicate"),
])
def test_validator_rejects_bad_data(tmp_path, assembled, sheet, key, field, value, expected):
    obs, _ = assembled
    errors = validate(read_observations(_mutated_copy(tmp_path, sheet, key, field, value)), obs)
    assert any(expected in e for e in errors), errors


def test_workbook_formulas_recalculate_independently():
    from pipeline.verify_workbook import verify
    n, problems = verify(WORKBOOK)
    assert n > 1000 and problems == []


def test_kpi_sheets_match_the_python_results(assembled):
    _, results = assembled
    errors, compared = compare_with_workbook_formulas(results)
    assert errors == [] and compared > 800


def test_site_json_matches_the_dataset(assembled, site_data):
    obs, results = assembled
    for o in obs:
        got = site_data["observations"][o["metric"]["id"]][o["period"]]
        assert got["value"] == (None if o["value"] is None else float(o["value"])), o["key"]
        assert got["status"] == o["status"] and got["page"] == o["page"], o["key"]
    for kid, k in site_data["kpis"].items():
        for p in dataset.PERIODS_A + dataset.PERIODS_Q:
            want = results[p][kid]
            assert k["values"][p] == (None if want is None else float(want)), (kid, p)
    # The key findings and section notes are written once, from the figures, for all three versions.
    assert site_data["findings"] == findings(obs, results)
    assert site_data["notes"] == section_notes(obs, results)


def test_site_download_is_the_handover_workbook():
    assert SITE_DOWNLOAD.read_bytes() == WORKBOOK.read_bytes()


# ---------------------------------------------------------------- Excel dashboard
@pytest.mark.parametrize("module", ["build_excel_dashboard", "build_powerbi"])
def test_wrapped_text_keeps_units_and_periods_together(module):
    """Excel and Power BI wrap text at any space or after any hyphen, so the builders glue figures to their units and
    reports to their periods with non-breaking characters (an editor once turned them into plain spaces)."""
    import importlib
    glue = importlib.import_module(f"pipeline.{module}").glue
    assert glue("4,6 tỷ USD") == "4,6 tỷ USD"
    assert glue("+7,1 điểm %") == "+7,1 điểm %"
    assert glue("152 ngày") == "152 ngày"
    assert glue("10-Q Q3 FY2020") == "10‑Q Q3 FY2020"
    assert glue("FY2019 → FY2021") == "FY2019 → FY2021"
    assert glue("non-GAAP") == "non‑GAAP"


def _data_blocks(ws):
    """Blocks of the hidden chart-data sheet: {block name: {column header: [values]}}."""
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    blocks, i = {}, 0
    while i < len(rows) - 1:
        if rows[i][0] and rows[i + 1][0] in ("Kỳ", "Năm"):
            head = [h for h in rows[i + 1] if h is not None]
            j = i + 2
            while j < len(rows) and any(v is not None for v in rows[j][:len(head)]):
                j += 1
            blocks[rows[i][0]] = {h: [r[k] for r in rows[i + 2:j]] for k, h in enumerate(head)}
            i = j
        else:
            i += 1
    return blocks


def _check_excel_dashboard(path, site_data):
    from pipeline.build_excel_dashboard import (CHANNEL_ROWS, FINANCE_ROWS, SHEET_DATA, SHEET_OVERVIEW, SHEET_TABLES, SHEETS,
                                                glue, region_rows)
    wb = load_workbook(path)
    assert wb.sheetnames == SHEETS + [SHEET_DATA]
    assert wb[SHEET_DATA].sheet_state == "hidden"
    A_, Q_ = site_data["periods"]["A"], site_data["periods"]["Q"]
    k = lambda kid, ps: [site_data["kpis"][kid]["values"][p] for p in ps]
    bn = lambda mid, ps: [site_data["observations"][mid][p]["value"] / 1000 for p in ps]
    blocks = _data_blocks(wb[SHEET_DATA])
    # Excel keeps 15 to 17 significant digits when it saves, so numbers compare with a relative tolerance.
    for block, column, want in (("Cơ cấu kênh", "NIKE Direct", k("direct_share", A_)),
                                ("Tỷ trọng NIKE Direct theo quý", "Tỷ trọng", k("direct_share", Q_)),
                                ("Doanh thu theo kênh, từng quý", "NIKE Direct", bn("nike_direct_revenue", Q_)),
                                ("Doanh thu theo kênh, từng quý", "Bán buôn", bn("nike_brand_wholesale_revenue", Q_)),
                                ("Digital / NIKE Direct, từng quý", "Tỷ trọng Digital", k("digital_share_of_direct", Q_)),
                                ("Biên gộp và biên EBIT, từng quý", "Biên gộp", k("gross_margin", Q_)),
                                ("Biên gộp và biên EBIT, từng quý", "Biên EBIT", k("ebit_margin", Q_)),
                                ("Tồn kho cuối quý", "Tồn kho (tỷ USD)", bn("inventories", Q_)),
                                ("Số ngày tồn kho, từng quý", "Số ngày", k("days_inventory", Q_))):
        assert blocks[block][column] == pytest.approx(want, rel=1e-12), (block, column)
    texts = {c.value for row in wb[SHEET_OVERVIEW].iter_rows() for c in row if isinstance(c.value, str)}
    for f in site_data["findings"]:
        assert glue(f["value"]) in texts and glue(f["text"]) in texts, f["id"]
    # Tables: a reported figure has its page in the small row below it, so every figure ends under its header.
    ws, page_rows = wb[SHEET_TABLES], 0
    for row in ws.iter_rows(min_row=5, max_col=15, values_only=True):
        cells = [v or "" for v in row[3:]]
        if row[1] is None and any(cells) and not any(re.fullmatch(r"FY\d{4}", v) for v in cells):
            assert all(re.fullmatch(r"(† )?(tr\. \d+)?", v) for v in cells), row
            page_rows += 1
        else:
            assert not any("tr." in v for v in cells), row
    reported = lambda rows: sum(r[0] == "obs" for r in rows)
    # The annual and quarterly tables share their rows; the regional table follows.
    assert page_rows == 2 * reported(CHANNEL_ROWS + FINANCE_ROWS) + reported(region_rows(site_data["regions"])), page_rows
    # Every printed page of the tables opens with a table title or the column headers (a break id is the last row of a
    # page, 1-based).
    for brk in ws.row_breaks.brk:
        first = [ws.cell(brk.id + 1, c).value for c in (2, 4)]
        assert first[0] == "Chỉ tiêu" or first[1] == A_[0] or str(first[0]).startswith(("Theo quý", "Bốn khu vực")), (brk.id, first)
    _check_print_layout(path)


XL = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "x": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
      "p": "http://schemas.openxmlformats.org/package/2006/relationships"}
REL_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def _check_print_layout(path):
    """The chart sheets print at a fixed scale, with manual page breaks that never cut a chart in two; the tables
    sheet starts each table on a new page, and a page never ends between a figure and its page reference."""
    def part(base, target):     # relationship targets are absolute or relative to the part's folder
        return target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
    with zipfile.ZipFile(path) as z:
        rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
        sheets = [part("xl", rels[s.get(REL_ID)]) for s in ET.fromstring(z.read("xl/workbook.xml")).find("m:sheets", XL)]
        for name in sheets[:3]:
            ws = ET.fromstring(z.read(name))
            assert ws.find("m:pageSetup", XL).get("scale") == "79", name
            breaks = [int(b.get("id")) for b in ws.iterfind("m:rowBreaks/m:brk", XL)]
            srels = ET.fromstring(z.read(name.replace("worksheets/", "worksheets/_rels/") + ".rels"))
            drawing = next(part("xl/worksheets", r.get("Target")) for r in srels.iterfind("p:Relationship", XL)
                           if r.get("Type").endswith("/drawing"))
            anchors = list(ET.fromstring(z.read(drawing)).iter(f"{{{XL['x']}}}twoCellAnchor"))
            assert len(anchors) >= 2, drawing
            for a in anchors:
                top = int(a.find("x:from/x:row", XL).text)
                bottom = int(a.find("x:to/x:row", XL).text) - (a.find("x:to/x:rowOff", XL).text == "0")
                assert not [b for b in breaks if top < b <= bottom], (name, top, bottom, breaks)
        tables = ET.fromstring(z.read(sheets[3]))
        assert tables.find("m:pageSetup", XL).get("scale") == "66"
        assert len(list(tables.iterfind("m:rowBreaks/m:brk", XL))) == 5     # three tables of two pages each


def test_excel_dashboard_build_matches_site_data(tmp_path, site_data):
    from pipeline.build_excel_dashboard import build
    path = tmp_path / "dashboard.xlsx"
    build(path, save_with_excel=False)
    _check_excel_dashboard(path, site_data)


def test_delivered_excel_dashboard_is_current(site_data):
    """The handover copy is rebuilt from the current data and saved once by Excel, which stores the chart caches
    Excel for the web draws; its file properties carry the builder's name, not Excel's user."""
    from pipeline.build_excel_dashboard import OUT
    _check_excel_dashboard(OUT, site_data)
    with zipfile.ZipFile(OUT) as z:
        charts = [n for n in z.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", n)]
        assert len(charts) == 12 and all(b"xmlns:c16r2" in z.read(n) for n in charts)
        core = z.read("docProps/core.xml").decode("utf-8")
    creator = re.search(r"<dc:creator>(.*?)</dc:creator>", core).group(1)
    modified = re.search(r"<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", core).group(1)
    # A CI checkout has no git user name; the delivered copy was built where there is one.
    assert creator and modified == creator == (git_user_name() or creator)


# ---------------------------------------------------------------- Power BI
def _tmdl_rows(path):
    """Rows of the #table literal embedded in a TMDL partition, as dicts keyed by column name."""
    text = path.read_text(encoding="utf-8")
    names = re.findall(r'#"((?:[^"]|"")*)" = [\w.]+', text)
    token = re.compile(r'"((?:[^"]|"")*)"|(null)|(-?[\d.]+(?:[eE][-+]?\d+)?)')
    rows = []
    for line in re.findall(r"^\s*\{(.*)\},?\s*$", text, re.M):
        vals = [m.group(1).replace('""', '"') if m.group(1) is not None else None if m.group(2) else float(m.group(3))
                for m in token.finditer(line)]
        rows.append(dict(zip(names, vals)))
    return rows


def _same(got, want):
    return got is None if want is None else got == pytest.approx(want, rel=1e-12)


def _strings(obj):
    """Every string in a parsed JSON document (the text of the report's text boxes and titles)."""
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for v in obj.values():
            yield from _strings(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from _strings(v)


@pytest.fixture(scope="module")
def pbip(tmp_path_factory, site_data):
    from pipeline import build_powerbi as pb
    out = tmp_path_factory.mktemp("pbip")
    w = pb.Wording()
    pb.semantic_model(site_data, w, out)
    pb.report(site_data, w, out)
    return out


def test_powerbi_model_matches_site_data(pbip, site_data):
    from pipeline.build_powerbi import NAME, glue
    tables = pbip / f"{NAME}.SemanticModel" / "definition" / "tables"
    for name, freq in (("Nam", "A"), ("Quy", "Q")):
        periods = site_data["periods"][freq]
        rows = {r["Kỳ"]: r for r in _tmdl_rows(tables / f"{name}.tmdl")}
        assert list(rows) == [site_data["labels"][p] for p in periods], name
        for p in periods:
            row = rows[site_data["labels"][p]]
            for mid, by_period in site_data["observations"].items():
                assert _same(row[mid], (by_period.get(p) or {}).get("value")), (name, mid, p)
            for kid, k in site_data["kpis"].items():
                assert _same(row[kid], k["values"][p]), (name, kid, p)
    # Table texts: the annual and regional tables give each figure's page in a column of its own, linked to the report;
    # the quarterly table has no page columns (twelve columns of figures already).
    filings = {f["url"] for f in site_data["filings"].values()}
    for name, with_pages in (("BangNam", True), ("BangQuy", False), ("BangKhuVuc", True)):
        cells = _tmdl_rows(tables / f"{name}.tmdl")
        figures, pages = [c for c in cells if c["Loại"] == "Số"], [c for c in cells if c["Loại"] == "Trang"]
        assert figures and len(pages) == (len(figures) if with_pages else 0), name
        assert all("tr." not in c["Hiển thị"] and c["Liên kết"] is None for c in figures), name
        for c in pages:
            assert re.fullmatch(r"(†|† tr\. \d+|tr\. \d+)?", c["Hiển thị"]), c
            assert (c["Liên kết"] in filings) if "tr." in c["Hiển thị"] else c["Liên kết"] is None, c
        # A page column has a blank header of non-breaking spaces, a different number per period: plain spaces would be
        # trimmed by Power BI and the page columns of the three years would merge into one.
        heads = {c["Kỳ"]: c["Cột"] for c in pages}
        assert all(set(h) == {" "} for h in heads.values()) and len(set(heads.values())) == len(heads), (name, heads)
    show = {(c["Chỉ tiêu"], c["Kỳ"]): c["Hiển thị"] for c in _tmdl_rows(tables / "BangQuy.tmdl") if c["Loại"] == "Số"}
    assert show[(glue("NIKE Direct, triệu USD"), "FY2021Q2")] == "4.310"
    assert show[(glue("Doanh số Digital, tỷ USD"), "FY2020Q4")] == "≈1,7"
    assert show[(glue("Tỷ trọng NIKE Direct do NIKE nêu (xấp xỉ)"), "FY2020Q4")] == "Không công bố"


def test_powerbi_report_pages(pbip, site_data):
    from pipeline.build_powerbi import NAME, glue
    defn = pbip / f"{NAME}.Report" / "definition"
    assert json.loads((defn / "pages" / "pages.json").read_text(encoding="utf-8"))["pageOrder"] == PAGES
    visuals = {f"{f.parents[2].name}/{f.parent.name}": json.loads(f.read_text(encoding="utf-8"))
               for f in defn.glob("pages/*/visuals/*/visual.json")}
    # Only the regional EBIT chart has a filter: it leaves out the NIKE Brand row, which has no EBIT of its own.
    assert [n for n, v in visuals.items() if v.get("filterConfig")] == ["khu_vuc_loi_nhuan/khu_vuc_ebit"]
    strings = set(_strings(list(visuals.values())))
    for f in site_data["findings"]:
        assert glue(f["text"]) in strings, f["id"]


def test_delivered_pbix_is_current(site_data):
    """The .pbix handed to the team is saved from the project by Power BI Desktop, so it can lag behind a rebuild:
    it must have the five pages, no saved filter but the regional one, and the current findings."""
    from pipeline.build_powerbi import glue
    with zipfile.ZipFile(dataset.ROOT / "handover" / "NIKE_Dashboard_FY2019-2021.pbix") as z:
        assert json.loads(z.read("Report/definition/pages/pages.json"))["pageOrder"] == PAGES
        visuals = {"/".join(n.split("/")[-4::2]): json.loads(z.read(n)) for n in z.namelist() if n.endswith("/visual.json")}
        assert z.getinfo("DataModel").file_size > 100_000
    assert [n for n, v in visuals.items() if v.get("filterConfig")] == ["khu_vuc_loi_nhuan/khu_vuc_ebit"]
    strings = set(_strings(list(visuals.values())))
    for f in site_data["findings"]:
        assert glue(f["text"]) in strings, f["id"]
