"""Recalculate every workbook formula with an independent Excel engine and compare.

Usage:  python -m pipeline.verify_workbook [path]

The `formulas` package evaluates the Excel formulas from scratch (it ignores the cached
values written by the builder). Each result must equal the cached value, i.e. the Python
evaluation of the same catalogue expression, and no cell may evaluate to an Excel error.
"""
import math
import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook

from .build_workbook import WORKBOOK

ERRORS = ("#VALUE!", "#REF!", "#N/A", "#NAME?", "#DIV/0!", "#NUM!", "#NULL!")


def _scalar(v):
    v = getattr(v, "value", v)
    while hasattr(v, "tolist") and not isinstance(v, (str, bytes)):
        v = v.tolist()
        if isinstance(v, list):
            v = v[0]
    if isinstance(v, list):
        v = v[0]
    return v


def verify(path=WORKBOOK):
    import formulas
    warnings.filterwarnings("ignore")
    path = Path(path)
    formulas_wb = load_workbook(path)
    cached_wb = load_workbook(path, data_only=True)
    cells = [(ws.title, c.coordinate) for ws in formulas_wb for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and c.value.startswith("=")]
    model = formulas.ExcelModel().loads(str(path)).finish()
    solution = model.calculate()
    computed = {}
    for key, val in solution.items():
        if "!" in key and ":" not in key.split("!")[1]:
            sheet = key.split("]")[1].split("'")[0] if "]" in key else key.split("!")[0]
            computed[(sheet.upper(), key.split("!")[1].replace("$", ""))] = _scalar(val)
    problems = []
    for sheet, coord in cells:
        got = computed.get((sheet.upper(), coord))
        want = cached_wb[sheet][coord].value
        if got is None and want in ("", None):
            continue
        if isinstance(got, str) and got in ERRORS or str(got) in ERRORS:
            problems.append(f"{sheet}!{coord}: evaluates to {got}")
            continue
        if isinstance(want, (int, float)) and isinstance(got, (int, float)):
            if not math.isclose(float(got), float(want), rel_tol=1e-9, abs_tol=1e-9):
                problems.append(f"{sheet}!{coord}: engine {got} vs cached {want}")
        elif str(got) != str(want if want is not None else ""):
            problems.append(f"{sheet}!{coord}: engine {got!r} vs cached {want!r}")
    return len(cells), problems


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else WORKBOOK
    n, problems = verify(path)
    print(f"{n} formulas recalculated independently; {len(problems)} problems")
    for p in problems[:50]:
        print("  ", p)
    return 1 if problems or n == 0 else 0


if __name__ == "__main__":
    sys.exit(main())
