"""Validate the handover workbook and export the dashboard data layer.

Usage:  python -m pipeline.build_site_data

Reads the reported observations from sheets 02–04 of the workbook (the source of truth
students may annotate), validates them against the extraction, recomputes every KPI with
the shared catalogue formulas, and writes site/data/nike-financials.json. Any failed check
stops the build with a non-zero exit code; nothing is exported from data that failed.
"""
import datetime
import json
import math
import shutil
import sys
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from . import dataset
from .build_workbook import (HEADER_ROW, OBS_COLUMNS, OBS_SHEETS, SHEET_KPI_A, SHEET_KPI_Q, WORKBOOK, date_text, load_quotes,
                             obs_key, period_dates)
from .insights import findings, section_notes
from .metrics import (A, CHECKS, COMPARISONS, KPIS, Q, REGIONS, REPORTED, REPORTED_BY_ID, STATUS_CONFLICT, STATUS_HUMAN,
                      STATUS_MISSING, STATUSES, applies, location)
from .quotes import QUOTES_BY_ID
from .recheck_pages import EVIDENCE_PAGE
from .sources import all_filings, document_title, filing, sec_document_url

ROOT = Path(__file__).resolve().parent.parent
SITE_DATA = ROOT / "site" / "data" / "nike-financials.json"
SITE_DOWNLOAD = ROOT / "site" / "downloads" / WORKBOOK.name
EVIDENCE = ROOT / "handover" / "bang-chung-doi-chieu"
SITE_EVIDENCE = ROOT / "site" / "bang-chung-doi-chieu"
HEADERS = {field: h for h, field, _ in OBS_COLUMNS}
MAY_BE_NEGATIVE = {"interest_expense_net", "other_expense_net", "corporate_revenue", "net_income", "income_before_taxes",
                   "income_tax_expense", "ebit_reported"}


class ValidationError(Exception):
    pass


def read_observations(path=WORKBOOK):
    frames = []
    for sheet in OBS_SHEETS:
        df = pd.read_excel(path, sheet_name=sheet, header=HEADER_ROW, dtype={HEADERS["page"]: "Int64"})
        missing_cols = [h for h in HEADERS.values() if h not in df.columns]
        if missing_cols:
            raise ValidationError(f"{sheet}: missing columns {missing_cols}")
        df["sheet_name"] = sheet
        frames.append(df)
    return pd.concat(frames, ignore_index=True).rename(columns={v: k for k, v in HEADERS.items()})


def _num(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    return Decimal(str(v))


def _blank(v):
    return v is None or pd.isna(v) or not str(v).strip()


def _same(a, b):
    return (a is None and b is None) or (a is not None and b is not None and abs(a - b) < Decimal("0.0000001"))


def validate(df, expected):
    """Workbook rows against the observations the extraction gives; returns a list of errors."""
    errors = []
    want = {obs_key(o["period"], o["metric"]["id"]): o for o in expected}
    keys = df["key"].astype(str)
    dup = keys[keys.duplicated()].tolist()
    if dup:
        errors.append(f"Duplicate observation keys: {dup}")
    missing_rows = sorted(set(want) - set(keys))
    extra_rows = sorted(set(keys) - set(want))
    if missing_rows:
        errors.append(f"Observations missing from workbook: {missing_rows}")
    if extra_rows:
        errors.append(f"Unexpected observations (unknown metric or period): {extra_rows}")
    for r in df.itertuples():
        o = want.get(r.key)
        if o is None:
            continue
        m = o["metric"]
        where = f"{r.key} ({r.sheet_name})"
        if r.metric_id != m["id"] or r.period_label != dataset.period_label(o["period"]):
            errors.append(f"{where}: key does not match period and metric")
        if r.unit != m["unit"]:
            errors.append(f"{where}: unit '{r.unit}' should be '{m['unit']}'")
        if r.scope != m["scope"]:
            errors.append(f"{where}: reporting scope '{r.scope}' should be '{m['scope']}'")
        if r.status not in STATUSES:
            errors.append(f"{where}: unknown data status '{r.status}'")
        value = _num(r.value)
        if value is None and r.status != STATUS_MISSING:
            errors.append(f"{where}: value is empty but status is '{r.status}' (empty values must be marked missing)")
        if value is not None and r.status == STATUS_MISSING:
            errors.append(f"{where}: has a value but is marked missing")
        if r.status == STATUS_CONFLICT:
            errors.append(f"{where}: conflicting evidence must be resolved before publishing")
        if r.status == STATUS_HUMAN and (_blank(r.verified_by) or _blank(r.verified_on)):
            errors.append(f"{where}: human verification needs the verifier's name and the date")
        if not _same(value, o["value"]):
            errors.append(f"{where}: value {value} differs from the extraction ({o['value']}); correct the workbook or re-run "
                          f"extraction")
        if not _same(_num(r.prior), o["prior"]):
            errors.append(f"{where}: prior-period value {_num(r.prior)} differs from the report's comparative {o['prior']}")
        if value is not None:
            if value < 0 and m["id"] not in MAY_BE_NEGATIVE and m["unit"] != "%":
                errors.append(f"{where}: unexpected negative value {value}")
            if m["unit"] == "%" and abs(value) > 5:
                errors.append(f"{where}: percentage {value} looks like it was entered in whole percent instead of a fraction")
            if r.page is pd.NA or _blank(r.section) or str(r.section) == "—" or _blank(r.evidence) or _blank(r.accession):
                errors.append(f"{where}: incomplete provenance (page, section, evidence and accession are required)")
            if r.accession != o["accession"]:
                errors.append(f"{where}: accession {r.accession} is not the one of {o['source']}")
            if r.page is not pd.NA and o["page"] is not None and int(r.page) != o["page"]:
                errors.append(f"{where}: page {r.page} differs from the extraction (page {o['page']})")
    return errors


def observations_from(df, expected):
    """The extraction's observations with the team's status and verification columns from the workbook."""
    rows = {r.key: r for r in df.itertuples()}
    for o in expected:
        r = rows[obs_key(o["period"], o["metric"]["id"])]
        o.update(status=r.status, verified_by="" if pd.isna(r.verified_by) else str(r.verified_by),
                 verified_on="" if pd.isna(r.verified_on) else date_text(r.verified_on), cross_check=r.cross_check,
                 restatement=r.restatement)
    return expected


def check_results(results):
    errors = []
    for period, env in results.items():
        freq = dataset.freq_of(period)
        for c in CHECKS + COMPARISONS:
            v = env.get(c["id"])
            if v is None:
                # Comparisons need a figure NIKE states, which the fourth quarter does not have.
                if c in CHECKS and applies(c, freq):
                    errors.append(f"Check {c['id']} {period}: cannot be computed (missing input)")
                continue
            if abs(v) > Decimal(str(c["tol"])):
                errors.append(f"Check {c['id']} {period}: difference {v} exceeds tolerance {c['tol']}")
    return errors


def compare_with_workbook_formulas(results, path=WORKBOOK):
    """Cached results of the KPI sheet formulas must equal the Python results."""
    wb = load_workbook(path, data_only=True, read_only=True)
    errors, compared = [], 0
    points = {c["id"] for c in COMPARISONS}    # shown in percentage points (×100)
    for sheet, freq in ((SHEET_KPI_A, A), (SHEET_KPI_Q, Q)):
        rows = list(wb[sheet].iter_rows(min_row=HEADER_ROW + 1, values_only=True))
        header = rows[0]
        labels = {dataset.period_label(p): p for p in dataset.PERIODS[freq]}
        cols = {labels[h]: i for i, h in enumerate(header) if h in labels}
        ids = {k["id"] for k in KPIS + CHECKS + COMPARISONS if applies(k, freq)}
        for row in rows[1:]:
            if row[0] not in ids:
                continue
            for p, i in cols.items():
                want, got = results[p][row[0]], row[i]
                if want is not None and row[0] in points:
                    want *= 100
                compared += 1
                if want is None:
                    if got not in ("–", None):
                        errors.append(f"{sheet} {row[0]} {p}: expected '–', workbook shows {got!r}")
                elif not isinstance(got, (int, float)) or not math.isclose(float(want), got, rel_tol=1e-9, abs_tol=1e-9):
                    errors.append(f"{sheet} {row[0]} {p}: workbook {got!r} vs recomputed {float(want)}")
    if compared == 0:
        errors.append("KPI sheets: no KPI rows found to compare")
    return errors, compared


def _f(v):
    return None if v is None else float(v)


def export(obs, results, extracted_on):
    periods = dataset.PERIODS_A + dataset.PERIODS_Q
    sample = {o["period"]: o for o in obs if o["metric"]["id"] == "revenues"}
    quotes = []
    for q in load_quotes():
        f = filing(q["filing"])
        quotes.append(dict(id=q["id"], topic=QUOTES_BY_ID[q["id"]]["topic"], filing=q["filing"], page=int(q["page"]),
                           url=sec_document_url(f), text=q["text"], vi=QUOTES_BY_ID[q["id"]]["vi"]))
    data = {
        "meta": {
            "title": "NIKE: bán trực tiếp (D2C) và kết quả tài chính, FY2019–FY2021",
            "generated_on": datetime.date.today().isoformat(),
            "extracted_on": extracted_on,
            "workbook": f"downloads/{WORKBOOK.name}",
            "evidence": f"bang-chung-doi-chieu/{EVIDENCE_PAGE}",
            "currency": "USD",
            "fiscal_year_end": "31/5",
            "status_counts": {s: sum(1 for o in obs if o["status"] == s) for s in STATUSES},
            "status_labels": STATUSES,
        },
        "periods": {"A": dataset.PERIODS_A, "Q": dataset.PERIODS_Q},
        "labels": {p: dataset.period_label(p) for p in periods},
        "dates": {p: period_dates(sample[p]) for p in periods},
        "filings": {f["id"]: {"accession": f["accession"], "filed": f["filed"], "period_end": f["period_end"], "form": f["form"],
                              "role": f["role"], "url": sec_document_url(f), "title": document_title(f)} for f in all_filings()},
        "metrics": {m["id"]: {"vi": m["vi"], "en": m["en"], "unit": m["unit"], "scope": m["scope"], "channel": m["channel"],
                              "region": m["region"], "kind": m["kind"], "definition": m["definition"],
                              "location_10k": location(m, "10-K"), "location_10q": location(m, "10-Q"),
                              "freq": sorted(m["freq"])} for m in REPORTED},
        "observations": {m["id"]: {} for m in REPORTED},
        "kpis": {k["id"]: {"vi": k["vi"], "unit": k["unit"], "group": k["group"], "note": k["note"], "freq": sorted(k["freq"]),
                           "values": {p: _f(results[p][k["id"]]) for p in periods}} for k in KPIS},
        "checks": {c["id"]: {"vi": c["vi"], "tol": c["tol"], "values": {p: _f(results[p][c["id"]]) for p in periods}}
                   for c in CHECKS + COMPARISONS},
        "regions": [{"id": r[0], "vi": r[1], "short": r[2], "en": r[3]} for r in REGIONS],
        "quotes": quotes,
        "findings": findings(obs, results),
        "notes": section_notes(obs, results),
    }
    for o in obs:
        data["observations"][o["metric"]["id"]][o["period"]] = {
            "value": _f(o["value"]), "prior": _f(o["prior"]), "status": o["status"], "page": o["page"], "page2": o["page2"],
            "source": o["source"], "source2": o["source2"], "section": o["section"], "evidence": o["evidence"],
            "cross_check": o["cross_check"], "restatement": o["restatement"], "method": o["method"], "approx": o["approx"],
            "verified_by": o["verified_by"], "verified_on": o["verified_on"]}
    SITE_DATA.parent.mkdir(parents=True, exist_ok=True)
    SITE_DATA.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    SITE_DOWNLOAD.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(WORKBOOK, SITE_DOWNLOAD)
    if EVIDENCE.exists():   # the evidence page for the MD&A-only figures, so dashboard viewers can open it
        shutil.copytree(EVIDENCE, SITE_EVIDENCE, dirs_exist_ok=True)
    return data


def main():
    extracted_on = json.loads((ROOT / "data" / "extracted" / "meta.json").read_text())["extracted_on"]
    expected, _ = dataset.assemble()
    df = read_observations()
    errors = validate(df, expected)
    if errors:
        print("VALIDATION FAILED – dashboard data not exported:")
        for e in errors[:60]:
            print("  ✗", e)
        return 1
    obs = observations_from(df, expected)
    results = dataset.compute(obs)
    errors = check_results(results)
    wb_errors, compared = compare_with_workbook_formulas(results)
    errors += wb_errors
    if errors:
        print("VALIDATION FAILED – dashboard data not exported:")
        for e in errors[:60]:
            print("  ✗", e)
        return 1
    data = export(obs, results, extracted_on)
    n_values = sum(1 for o in obs if o["value"] is not None)
    print(f"✓ {len(obs)} observations validated ({n_values} with values, "
          f"{data['meta']['status_counts'][STATUS_MISSING]} marked missing)")
    print(f"✓ {len(CHECKS)} arithmetic reconciliations and {len(COMPARISONS)} comparisons with NIKE-reported figures pass "
          f"for {len(results)} periods")
    print(f"✓ {compared} KPI sheet formula results match the recomputed KPIs")
    print(f"✓ Wrote {SITE_DATA.relative_to(ROOT)}, {SITE_DOWNLOAD.relative_to(ROOT)} and {SITE_EVIDENCE.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
